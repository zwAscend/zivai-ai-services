from __future__ import annotations

import base64
import io
import json
import os
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Sequence

import pdfplumber
from docx import Document as DocxDocument
from PIL import Image
from werkzeug.datastructures import FileStorage
from werkzeug.utils import secure_filename

from .schema import OcrDocumentResult, OcrGeneralRequest, OcrGeneralResponse, OcrReferenceDocument


OCR_ENGINE = "huawei-ocr-general-text"
DIGITAL_TEXT_ENGINE = "builtin-digital-text"
IMAGE_EXTENSIONS = {"jpeg", "jpg", "png", "bmp", "gif", "tiff", "tif", "webp", "pcx", "ico", "psd"}
TEXT_EXTENSIONS = {"txt", "md", "csv", "json", "log"}
DOCX_EXTENSIONS = {"docx"}
PDF_EXTENSIONS = {"pdf"}
DEFAULT_ALLOWED_FORMATS = IMAGE_EXTENSIONS | TEXT_EXTENSIONS | DOCX_EXTENSIONS | PDF_EXTENSIONS


class OcrValidationError(ValueError):
    pass


class OcrConfigurationError(RuntimeError):
    pass


class OcrProcessingError(RuntimeError):
    pass


@dataclass(frozen=True)
class HuaweiOcrSettings:
    ak: str | None
    sk: str | None
    project_id: str | None
    ocr_endpoint: str
    host: str
    force_trailing_slash: bool
    http_timeout_seconds: int
    detect_direction: bool
    quick_mode: bool
    max_original_file_size_bytes: int
    max_encoded_image_bytes: int
    max_image_width: int
    max_image_height: int
    jpeg_quality: float
    min_jpeg_quality: float
    adaptive_resize_percent: int
    max_adaptive_passes: int
    pdf_render_dpi: int
    max_pdf_pages: int
    allowed_formats: tuple[str, ...]

    @classmethod
    def from_env(cls) -> "HuaweiOcrSettings":
        allowed = os.getenv("HWC_GENERAL_TEXT_ALLOWED_FORMATS", "")
        normalized_allowed = tuple(
            sorted(
                {
                    part.strip().lower().lstrip(".")
                    for part in allowed.split(",")
                    if part.strip()
                }
                or DEFAULT_ALLOWED_FORMATS
            )
        )
        return cls(
            ak=_trim_to_none(os.getenv("HWC_AK")),
            sk=_trim_to_none(os.getenv("HWC_SK")),
            project_id=_trim_to_none(os.getenv("HWC_PROJECT_ID")),
            ocr_endpoint=(
                _trim_to_none(os.getenv("HWC_OCR_ENDPOINT"))
                or "https://ocr.ap-southeast-1.myhuaweicloud.com"
            ),
            host=_trim_to_none(os.getenv("HWC_HOST")) or "ocr.ap-southeast-1.myhuaweicloud.com",
            force_trailing_slash=_env_bool("HWC_FORCE_TRAILING_SLASH", False),
            http_timeout_seconds=_env_int("HWC_HTTP_TIMEOUT_SECONDS", 180),
            detect_direction=_env_bool("HWC_GENERAL_TEXT_DETECT_DIRECTION", True),
            quick_mode=_env_bool("HWC_GENERAL_TEXT_QUICK_MODE", False),
            max_original_file_size_bytes=_env_int("HWC_GENERAL_TEXT_MAX_ORIGINAL_FILE_SIZE_BYTES", 7 * 1024 * 1024),
            max_encoded_image_bytes=_env_int("HWC_GENERAL_TEXT_MAX_ENCODED_IMAGE_BYTES", 2_500_000),
            max_image_width=_env_int("HWC_GENERAL_TEXT_MAX_IMAGE_WIDTH", 2200),
            max_image_height=_env_int("HWC_GENERAL_TEXT_MAX_IMAGE_HEIGHT", 2200),
            jpeg_quality=_env_float("HWC_GENERAL_TEXT_JPEG_QUALITY", 0.72),
            min_jpeg_quality=_env_float("HWC_GENERAL_TEXT_MIN_JPEG_QUALITY", 0.50),
            adaptive_resize_percent=_env_int("HWC_GENERAL_TEXT_ADAPTIVE_RESIZE_PERCENT", 85),
            max_adaptive_passes=_env_int("HWC_GENERAL_TEXT_MAX_ADAPTIVE_PASSES", 6),
            pdf_render_dpi=_env_int("HWC_GENERAL_TEXT_PDF_RENDER_DPI", 200),
            max_pdf_pages=_env_int("HWC_GENERAL_TEXT_MAX_PDF_PAGES", 50),
            allowed_formats=normalized_allowed,
        )


class _HuaweiOcrClient:
    def __init__(self, settings: HuaweiOcrSettings):
        self.settings = settings
        self.client = _build_huawei_client(
            settings.ak or "",
            settings.sk or "",
            settings.project_id or "",
            _resolve_endpoint(settings),
            settings.http_timeout_seconds,
        )

    def recognize_base64(
        self,
        image_b64: str,
        *,
        request_options: OcrGeneralRequest,
        pdf_page_number: int | None = None,
    ) -> dict:
        from huaweicloudsdkocr.v1.model.general_text_request_body import GeneralTextRequestBody
        from huaweicloudsdkocr.v1.model.recognize_general_text_request import RecognizeGeneralTextRequest

        body = GeneralTextRequestBody(
            image=image_b64,
            detect_direction=request_options.detectDirection,
            quick_mode=request_options.quickMode,
            character_mode=request_options.characterMode,
            language=request_options.language,
            single_orientation_mode=request_options.singleOrientationMode,
            pdf_page_number=pdf_page_number,
            return_markdown_result=request_options.returnMarkdownResult,
        )
        sdk_request = RecognizeGeneralTextRequest(body=body)
        response = self.client.recognize_general_text(sdk_request)
        if hasattr(response, "to_dict"):
            return response.to_dict()
        if hasattr(response, "to_str"):
            return json.loads(response.to_str())
        return json.loads(json.dumps(response, default=str))


@lru_cache(maxsize=4)
def _build_huawei_client(ak: str, sk: str, project_id: str, endpoint: str, timeout_seconds: int):
    from huaweicloudsdkcore.auth.credentials import BasicCredentials
    from huaweicloudsdkcore.http.http_config import HttpConfig
    from huaweicloudsdkocr.v1 import OcrClient

    credentials = BasicCredentials(ak, sk, project_id)
    http_config = HttpConfig.get_default_config()
    http_config.timeout = timeout_seconds
    return OcrClient.new_builder().with_credentials(credentials).with_http_config(http_config).with_endpoint(endpoint).build()


def extract_general_ocr(
    files: Sequence[FileStorage],
    request_options: OcrGeneralRequest,
) -> OcrGeneralResponse:
    settings = HuaweiOcrSettings.from_env()
    normalized_files = [item for item in files if item and getattr(item, "filename", None)]

    if not normalized_files and not request_options.url:
        raise OcrValidationError("Provide at least one file, or supply a url in the OCR request.")

    warnings: list[str] = []
    documents: list[OcrDocumentResult] = []

    for uploaded in normalized_files:
        documents.extend(_process_uploaded_file(uploaded, request_options, settings, warnings))

    if request_options.url:
        documents.append(_process_remote_url(request_options, settings))

    combined_text = "\n\n".join(
        item.markdown.strip() or item.fullText.strip()
        for item in documents
        if (item.markdown.strip() or item.fullText.strip())
    ).strip()
    reference_documents = [item.referenceDocument for item in documents if item.referenceDocument.markdown.strip()]

    return OcrGeneralResponse(
        request=request_options,
        documents=documents,
        combinedText=combined_text,
        referenceDocuments=reference_documents,
        warnings=warnings,
    )


def _process_uploaded_file(
    uploaded: FileStorage,
    request_options: OcrGeneralRequest,
    settings: HuaweiOcrSettings,
    warnings: list[str],
) -> list[OcrDocumentResult]:
    raw_bytes = uploaded.read()
    filename = secure_filename(uploaded.filename or "upload") or "upload"
    extension = _resolve_extension(filename, uploaded.mimetype)
    if extension not in settings.allowed_formats:
        raise OcrValidationError(
            f"Unsupported file format '{extension}'. Allowed formats: {', '.join(settings.allowed_formats)}"
        )

    if settings.max_original_file_size_bytes > 0 and len(raw_bytes) > settings.max_original_file_size_bytes:
        raise OcrValidationError(
            f"File '{filename}' exceeds the configured max size of {settings.max_original_file_size_bytes} bytes."
        )

    if extension in TEXT_EXTENSIONS:
        return [_extract_text_document(filename, extension, raw_bytes, request_options)]
    if extension in DOCX_EXTENSIONS:
        return [_extract_docx_document(filename, raw_bytes, request_options)]
    if extension in PDF_EXTENSIONS:
        return _extract_pdf_document(filename, raw_bytes, request_options, settings, warnings)
    return [_extract_image_document(filename, extension, raw_bytes, request_options, settings)]


def _extract_text_document(filename: str, extension: str, raw_bytes: bytes, request_options: OcrGeneralRequest) -> OcrDocumentResult:
    text = raw_bytes.decode("utf-8", errors="ignore").strip()
    if not text:
        raise OcrValidationError(f"File '{filename}' does not contain usable text.")
    return _build_result(
        document_name=filename,
        file_format=extension,
        page_number=1,
        total_pages=1,
        engine=DIGITAL_TEXT_ENGINE,
        mode="digital_text",
        full_text=text,
        markdown=text,
        request_options=request_options,
        average_confidence=None,
        words_block_count=max(1, len(text.splitlines())),
        metadata={"source": "text-file"},
    )


def _extract_docx_document(filename: str, raw_bytes: bytes, request_options: OcrGeneralRequest) -> OcrDocumentResult:
    with io.BytesIO(raw_bytes) as handle:
        document = DocxDocument(handle)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    text = "\n".join(paragraphs).strip()
    if not text:
        raise OcrValidationError(f"DOCX file '{filename}' does not contain usable text.")
    return _build_result(
        document_name=filename,
        file_format="docx",
        page_number=1,
        total_pages=1,
        engine=DIGITAL_TEXT_ENGINE,
        mode="digital_text",
        full_text=text,
        markdown=text,
        request_options=request_options,
        average_confidence=None,
        words_block_count=max(1, len(paragraphs)),
        metadata={"source": "docx"},
    )


def _extract_pdf_document(
    filename: str,
    raw_bytes: bytes,
    request_options: OcrGeneralRequest,
    settings: HuaweiOcrSettings,
    warnings: list[str],
) -> list[OcrDocumentResult]:
    digital_pages = _extract_pdf_text_pages(filename, raw_bytes, request_options, settings)
    if digital_pages and request_options.preferDigitalExtraction and not request_options.forceOcr:
        warnings.append(
            f"Used built-in digital PDF extraction for '{filename}' where selectable text was available."
        )
        return digital_pages

    warnings.append(f"Using Huawei OCR for PDF '{filename}'.")
    return _ocr_pdf_pages(filename, raw_bytes, request_options, settings)


def _extract_pdf_text_pages(
    filename: str,
    raw_bytes: bytes,
    request_options: OcrGeneralRequest,
    settings: HuaweiOcrSettings,
) -> list[OcrDocumentResult]:
    results: list[OcrDocumentResult] = []
    with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
        total_pages = len(pdf.pages)
        if total_pages == 0:
            raise OcrValidationError(f"PDF '{filename}' has no pages.")
        max_pages = _effective_page_limit(total_pages, request_options.maxPages, settings.max_pdf_pages)
        for index, page in enumerate(pdf.pages[:max_pages], start=1):
            text = (page.extract_text() or "").strip()
            if not text:
                continue
            results.append(
                _build_result(
                    document_name=filename,
                    file_format="pdf",
                    page_number=index,
                    total_pages=total_pages,
                    engine=DIGITAL_TEXT_ENGINE,
                    mode="digital_text",
                    full_text=text,
                    markdown=text,
                    request_options=request_options,
                    average_confidence=None,
                    words_block_count=max(1, len(text.splitlines())),
                    metadata={"source": "pdfplumber"},
                )
            )
    return results


def _ocr_pdf_pages(
    filename: str,
    raw_bytes: bytes,
    request_options: OcrGeneralRequest,
    settings: HuaweiOcrSettings,
) -> list[OcrDocumentResult]:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(io.BytesIO(raw_bytes))
    total_pages = len(pdf)
    if total_pages == 0:
        raise OcrValidationError(f"PDF '{filename}' has no pages.")

    max_pages = _effective_page_limit(total_pages, request_options.maxPages, settings.max_pdf_pages)
    results: list[OcrDocumentResult] = []
    scale = max(1.0, settings.pdf_render_dpi / 72.0)

    for page_number in range(1, max_pages + 1):
        page = pdf[page_number - 1]
        image = page.render(scale=scale).to_pil()
        page_bytes = _encode_image_with_limits(image, settings)
        raw = _recognize_with_huawei(
            image_bytes=page_bytes,
            request_options=request_options,
            settings=settings,
            pdf_page_number=page_number,
        )
        results.append(
            _map_huawei_result(
                raw=raw,
                document_name=filename,
                file_format="pdf",
                page_number=page_number,
                total_pages=total_pages,
                request_options=request_options,
            )
        )
    return results


def _extract_image_document(
    filename: str,
    extension: str,
    raw_bytes: bytes,
    request_options: OcrGeneralRequest,
    settings: HuaweiOcrSettings,
) -> OcrDocumentResult:
    raw = _recognize_with_huawei(
        image_bytes=raw_bytes,
        request_options=request_options,
        settings=settings,
        pdf_page_number=None,
    )
    return _map_huawei_result(
        raw=raw,
        document_name=filename,
        file_format=extension,
        page_number=1,
        total_pages=1,
        request_options=request_options,
    )


def _process_remote_url(request_options: OcrGeneralRequest, settings: HuaweiOcrSettings) -> OcrDocumentResult:
    if not _has_huawei_credentials(settings):
        raise OcrConfigurationError(
            "Huawei OCR credentials are required for URL-based OCR. Set HWC_AK, HWC_SK, and HWC_PROJECT_ID."
        )
    client = _HuaweiOcrClient(settings)
    from huaweicloudsdkocr.v1.model.general_text_request_body import GeneralTextRequestBody
    from huaweicloudsdkocr.v1.model.recognize_general_text_request import RecognizeGeneralTextRequest

    body = GeneralTextRequestBody(
        url=request_options.url,
        detect_direction=request_options.detectDirection,
        quick_mode=request_options.quickMode,
        character_mode=request_options.characterMode,
        language=request_options.language,
        single_orientation_mode=request_options.singleOrientationMode,
        return_markdown_result=request_options.returnMarkdownResult,
    )
    sdk_request = RecognizeGeneralTextRequest(body=body)
    response = client.client.recognize_general_text(sdk_request)
    raw = response.to_dict() if hasattr(response, "to_dict") else json.loads(response.to_str())
    document_name = Path(request_options.url or "remote").name or "remote"
    return _map_huawei_result(
        raw=raw,
        document_name=document_name,
        file_format=_resolve_extension(document_name, None),
        page_number=1,
        total_pages=1,
        request_options=request_options,
    )


def _recognize_with_huawei(
    *,
    image_bytes: bytes,
    request_options: OcrGeneralRequest,
    settings: HuaweiOcrSettings,
    pdf_page_number: int | None,
) -> dict:
    if not _has_huawei_credentials(settings):
        raise OcrConfigurationError(
            "Huawei OCR credentials are not configured. Set HWC_AK, HWC_SK, and HWC_PROJECT_ID to use OCR on images or scanned PDFs."
        )

    prepared = _prepare_image_payload(image_bytes, settings)
    client = _HuaweiOcrClient(settings)
    try:
        return client.recognize_base64(prepared, request_options=request_options, pdf_page_number=pdf_page_number)
    except Exception as exc:
        message = str(exc).strip() or exc.__class__.__name__
        raise OcrProcessingError(f"Huawei OCR request failed: {message}") from exc


def _prepare_image_payload(image_bytes: bytes, settings: HuaweiOcrSettings) -> str:
    image = Image.open(io.BytesIO(image_bytes))
    return base64.b64encode(_encode_image_with_limits(image, settings)).decode("utf-8")


def _encode_image_with_limits(image: Image.Image, settings: HuaweiOcrSettings) -> bytes:
    working = _normalize_image(image, settings)
    quality = max(settings.min_jpeg_quality, min(settings.jpeg_quality, 0.95))
    resize_ratio = max(0.5, min(settings.adaptive_resize_percent / 100.0, 0.95))

    for _attempt in range(max(1, settings.max_adaptive_passes)):
        payload = _encode_jpeg(working, quality)
        encoded_size = len(base64.b64encode(payload))
        if encoded_size <= settings.max_encoded_image_bytes:
            return payload
        quality = max(settings.min_jpeg_quality, quality - 0.08)
        working = _resize_by_ratio(working, resize_ratio)

    payload = _encode_jpeg(working, settings.min_jpeg_quality)
    encoded_size = len(base64.b64encode(payload))
    if encoded_size > settings.max_encoded_image_bytes:
        raise OcrValidationError(
            "Image remains too large for Huawei OCR after adaptive compression. Reduce dimensions or upload a clearer crop."
        )
    return payload


def _normalize_image(image: Image.Image, settings: HuaweiOcrSettings) -> Image.Image:
    working = image.convert("RGB") if image.mode not in {"RGB", "L"} else image
    width, height = working.size
    scale = min(
        1.0,
        settings.max_image_width / max(width, 1),
        settings.max_image_height / max(height, 1),
    )
    if scale >= 1.0:
        return working
    target = (
        max(1, int(round(width * scale))),
        max(1, int(round(height * scale))),
    )
    return working.resize(target, Image.Resampling.LANCZOS)


def _resize_by_ratio(image: Image.Image, ratio: float) -> Image.Image:
    width, height = image.size
    target = (max(1, int(round(width * ratio))), max(1, int(round(height * ratio))))
    if target == image.size:
        return image
    return image.resize(target, Image.Resampling.LANCZOS)


def _encode_jpeg(image: Image.Image, quality_ratio: float) -> bytes:
    quality = max(35, min(int(round(quality_ratio * 100)), 95))
    buffer = io.BytesIO()
    image.convert("RGB").save(buffer, format="JPEG", quality=quality, optimize=True)
    return buffer.getvalue()


def _map_huawei_result(
    *,
    raw: dict,
    document_name: str,
    file_format: str,
    page_number: int,
    total_pages: int,
    request_options: OcrGeneralRequest,
) -> OcrDocumentResult:
    result_node = raw.get("result") or {}
    words = result_node.get("words_block_list") or result_node.get("wordsBlockList") or []
    markdown = (result_node.get("markdown_result") or result_node.get("markdownResult") or "").strip()
    if markdown:
        full_text = markdown
    else:
        full_text = "\n".join(
            (block.get("words") or "").strip()
            for block in words
            if isinstance(block, dict) and (block.get("words") or "").strip()
        ).strip()

    confidences = [
        float(block.get("confidence"))
        for block in words
        if isinstance(block, dict) and block.get("confidence") is not None
    ]
    average_confidence = sum(confidences) / len(confidences) if confidences else None
    words_block_count = result_node.get("words_block_count") or result_node.get("wordsBlockCount") or len(words)

    return _build_result(
        document_name=document_name,
        file_format=file_format,
        page_number=page_number,
        total_pages=total_pages,
        engine=OCR_ENGINE,
        mode="ocr",
        full_text=full_text,
        markdown=markdown or full_text,
        request_options=request_options,
        average_confidence=average_confidence,
        words_block_count=int(words_block_count or 0),
        metadata={"rawResult": result_node},
    )


def _build_result(
    *,
    document_name: str,
    file_format: str,
    page_number: int,
    total_pages: int,
    engine: str,
    mode: str,
    full_text: str,
    markdown: str,
    request_options: OcrGeneralRequest,
    average_confidence: float | None,
    words_block_count: int,
    metadata: dict,
) -> OcrDocumentResult:
    normalized_markdown = markdown.strip() or full_text.strip()
    if total_pages > 1:
        ref_name = f"{document_name}#page-{page_number}"
    else:
        ref_name = document_name
    reference_document = OcrReferenceDocument(documentName=ref_name, markdown=normalized_markdown)
    return OcrDocumentResult(
        documentName=document_name,
        fileFormat=file_format,
        pageNumber=page_number,
        totalPages=total_pages,
        engine=engine,
        mode=mode,
        fullText=full_text.strip(),
        markdown=normalized_markdown,
        averageConfidence=average_confidence,
        wordsBlockCount=words_block_count,
        module=request_options.module,
        source=request_options.source,
        referenceDocument=reference_document,
        metadata=metadata,
    )


def _effective_page_limit(total_pages: int, request_limit: int | None, setting_limit: int) -> int:
    limit = total_pages
    if request_limit is not None:
        limit = min(limit, request_limit)
    if setting_limit > 0:
        limit = min(limit, setting_limit)
    return max(1, limit)


def _resolve_extension(filename: str, mimetype: str | None) -> str:
    suffix = Path(filename).suffix.lower().lstrip(".")
    if suffix:
        return suffix
    if mimetype == "application/pdf":
        return "pdf"
    if mimetype and mimetype.startswith("image/"):
        return mimetype.split("/", 1)[1].lower()
    if mimetype == "text/plain":
        return "txt"
    return "bin"


def _resolve_endpoint(settings: HuaweiOcrSettings) -> str:
    endpoint = (settings.ocr_endpoint or f"https://{settings.host}").strip()
    if settings.force_trailing_slash and not endpoint.endswith("/"):
        endpoint = endpoint + "/"
    while endpoint.endswith("/"):
        endpoint = endpoint[:-1]
    return endpoint


def _has_huawei_credentials(settings: HuaweiOcrSettings) -> bool:
    return bool(settings.ak and settings.sk and settings.project_id)


def _trim_to_none(value: str | None) -> str | None:
    if value is None:
        return None
    text = value.strip()
    return text or None


def _env_bool(name: str, default: bool) -> bool:
    raw = os.getenv(name)
    if raw is None:
        return default
    return raw.strip().lower() in {"1", "true", "yes", "on"}


def _env_int(name: str, default: int) -> int:
    raw = os.getenv(name)
    if raw is None:
        return default
    try:
        return int(raw.strip())
    except ValueError:
        return default


def _env_float(name: str, default: float) -> float:
    raw = os.getenv(name)
    if raw is None:
        return default
    try:
        return float(raw.strip())
    except ValueError:
        return default
